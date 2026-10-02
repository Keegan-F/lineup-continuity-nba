"""
continuity_lib.py — Core library for "The Value of Lineup Continuity in the NBA"

SINGLE SOURCE OF TRUTH: the five canonical July-1 roster-tag workbooks
    /mnt/user-data/outputs/NBA_RosterTags_5ML_{season}.xlsx   (2021-22 .. 2025-26)

NEVER use any *_base_v2.xlsx / *RosterSpotsTags* / *_FIXED* / *CORRECTED_master* file.
Those are quarantined in /mnt/user-data/outputs/_DEPRECATED_base_v2/ because they lost
injury attribution and mis-tagged injured/traded starters (see HANDOFF.md, Bug #1).

Workbook structure (per season file):
  'Roster Tags'   : one row per player-season. Cols include Player-Team, Team, Roster Tag,
                    GP, GS. Tag in {ST, SB, DB, GT}.
  '<TEAM>'        : lineup detail. col2=MIN, col5=NetRtg, cols13-17=key1..key5 (player keys),
                    cols18-22=t1..t5 (tags), col23-26=nST/nSB/nDB/nGT, col27=Scenario.
  '<TEAM> Summary': pre-aggregated scenario rollup.
  '<TEAM> Grid'   : game-by-game grid; rows 'Starter 1'..'Starter 5' x 82 game columns.

Six scenarios (col27):
  Starter Baseline / Starters + Short Bench Player / Starters + Deep Bench Player /
  Short Bench Preserved / Short Bench Pilfered / Garbage Time

PRESERVE bundle = {Starters + Deep Bench Player, Short Bench Preserved}
PILFER   bundle = {Starters + Short Bench Player, Short Bench Pilfered}
NRG (Net Ratings Gap) = minute-weighted preserve net - minute-weighted pilfer net
"""
import re
import os
import glob
import unicodedata
from collections import defaultdict

import openpyxl

SEASONS = ['2021-22', '2022-23', '2023-24', '2024-25', '2025-26']
POSITIONS = ['PG', 'SG', 'SF', 'PF', 'C']
TAGS = ('ST', 'SB', 'DB', 'GT')


# ---------------------------------------------------------------- data location
# IMPORTANT: /mnt/user-data/outputs/ does NOT persist across conversations. Only files
# the user re-uploads appear (in /mnt/user-data/uploads/). So we DISCOVER the workbooks
# rather than hardcoding a path. Set CONTINUITY_DATA to override.
_SEARCH_DIRS = [
    os.environ.get('CONTINUITY_DATA', ''),
    './data', '../data', '/mnt/user-data/uploads', '/mnt/user-data/outputs',
    '/mnt/user-data/uploads/data', '/home/claude/data',
]


def _find(filename):
    """Locate a source workbook across the plausible directories."""
    for d in _SEARCH_DIRS:
        if not d:
            continue
        p = os.path.join(d, filename)
        if os.path.exists(p):
            return p
    # last resort: recursive search under the mounted dirs
    for root in ('/mnt/user-data', '/home/claude'):
        hits = glob.glob(os.path.join(root, '**', filename), recursive=True)
        # never return a quarantined file
        hits = [h for h in hits if '_DEPRECATED' not in h]
        if hits:
            return hits[0]
    raise FileNotFoundError(
        '%s not found. Searched: %s\n'
        'The five NBA_RosterTags_5ML_{season}.xlsx and five NBA_Player_PerGame_{season}_byTeam.xlsx '
        'files must be present. Upload them, or set CONTINUITY_DATA=/path/to/data.'
        % (filename, [d for d in _SEARCH_DIRS if d])
    )


def tags_path(season):
    """The canonical July-1 roster-tag workbook. NEVER _base_v2/_FIXED/CORRECTED_master."""
    return _find('NBA_RosterTags_5ML_%s.xlsx' % season)


def pergame_path(season):
    return _find('NBA_Player_PerGame_%s_byTeam.xlsx' % season)


def preflight():
    """Verify all ten required workbooks are reachable BEFORE doing any work.
    Call this first in any new session."""
    missing = []
    for s in SEASONS:
        for fn in ('NBA_RosterTags_5ML_%s.xlsx' % s, 'NBA_Player_PerGame_%s_byTeam.xlsx' % s):
            try:
                _find(fn)
            except FileNotFoundError:
                missing.append(fn)
    if missing:
        raise FileNotFoundError('Missing %d required workbook(s):\n  %s'
                                % (len(missing), '\n  '.join(missing)))
    print('preflight OK — all 10 source workbooks located')
    for s in SEASONS:
        print('   %s  tags=%s' % (s, tags_path(s)))
    return True

SCENARIO_KEY = {
    'Starter Baseline': 'base',
    'Starters + Short Bench Player': 'st_short',   # pilfer side (4S + 1 Short Bench)
    'Starters + Deep Bench Player': 'st_deep',     # preserve side (4S + 1 Deep Bench)
    'Short Bench Preserved': 'prsv',               # preserve side (bench)
    'Short Bench Pilfered': 'pilf',                # pilfer side (bench)
    'Garbage Time': 'gt',
}
PRESERVE_SCENARIOS = {'st_deep', 'prsv'}
PILFER_SCENARIOS = {'st_short', 'pilf'}

# Basketball-Reference per-game files use different codes than the roster tags.
# BUG #3: omitting this silently DROPS Brooklyn, Charlotte and Phoenix from any
# position-based analysis. Always map.
ABBR_FIX = {'BRK': 'BKN', 'CHO': 'CHA', 'PHO': 'PHX'}


# ---------------------------------------------------------------- name handling
def norm(s):
    """ASCII-fold + lowercase. Handles Doncic/Jokic/Bogdanovic etc."""
    return unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().strip().lower()


def strip_suffix(k):
    """Drop Jr./Sr./II/III/IV. BUG #4: 'D. Jones' (tags) vs 'D. Jones Jr.' (per-game)."""
    return re.sub(r'\s+(jr|sr|ii|iii|iv)\.?$', '', k).strip()


def initial_key(full_name):
    """'Derrick Jones Jr.' -> ('d', 'jones jr.')  — first initial + surname, for
    distinct-Starter matching against the Grid sheets."""
    parts = str(full_name).split()
    if not parts:
        return None
    return (norm(parts[0])[0] if parts[0] else '', norm(' '.join(parts[1:])))


def strip_team(key, team):
    """'D. Garland-LAC' -> 'd. garland'"""
    k = str(key)
    suffix = '-' + str(team)
    if k.endswith(suffix):
        k = k[:-len(suffix)]
    return norm(k)


# ---------------------------------------------------------------- loaders
def load_tags(season):
    """-> {team: {normname: tag}}, {team: {normname: (gp, gs)}}"""
    wb = openpyxl.load_workbook(tags_path(season), read_only=True, data_only=True)
    rows = list(wb['Roster Tags'].iter_rows(values_only=True))
    hdr = [str(c) for c in rows[0]]
    ix = {c: i for i, c in enumerate(hdr)}
    tags = defaultdict(dict)
    games = defaultdict(dict)
    for r in rows[1:]:
        tag = r[ix['Roster Tag']]
        if not r[ix['Player-Team']] or tag not in TAGS:
            continue
        team = r[ix['Team']]
        nm = strip_team(r[ix['Player-Team']], team)
        tags[team][nm] = tag
        gp = r[ix['GP']] if isinstance(r[ix['GP']], (int, float)) else 0
        gs = r[ix['GS']] if isinstance(r[ix['GS']], (int, float)) else 0
        games[team][nm] = (gp, gs)
    wb.close()
    return dict(tags), dict(games)


def tag_key_forms(full_name):
    """BUG #5. The roster tags EXTEND the first initial to disambiguate teammates:
    'Jalen Williams' -> 'jal. williams', 'Jaylin Williams' -> 'jay. williams',
    'Stephen Curry' -> 'ste. curry', 'Jeff Green' -> 'jef. green'.

    load_positions() used to collapse the per-game full name down to 'j. williams',
    which (a) never matched the extended tag key -- so Jalen Williams, Jalen Green and
    STEPHEN CURRY had no position and were silently dropped from Finding 3.9 -- and
    (b) collided both Williamses onto one key, so one player's position overwrote the
    other's.

    The per-game file has the FULL name. Use it. Candidate keys, longest prefix first.
    """
    p = norm(full_name).split()
    if len(p) < 2:
        return [norm(full_name)]
    first, last = p[0], ' '.join(p[1:])
    forms = ['%s. %s' % (first[:k], last) for k in (4, 3, 2, 1) if len(first) >= k]
    out = []
    for f in forms:
        for cand in (f, strip_suffix(f)):
            if cand not in out:
                out.append(cand)
    return out


def load_positions(season, tags=None):
    """-> {team: {tagkey: POS}}

    If `tags` is supplied, per-game full names are resolved against the ACTUAL roster-tag
    keys (longest prefix wins). This is the bug #5 fix. Without `tags` it falls back to
    the old initial-collapsing behaviour, which is retained only so the old numbers can
    still be reproduced for comparison.
    """
    wb = openpyxl.load_workbook(pergame_path(season), read_only=True, data_only=True)
    posmap = defaultdict(dict)
    collisions = []
    for sheet in wb.sheetnames:
        team = ABBR_FIX.get(sheet, sheet)
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        hdr = [str(c) for c in rows[0]]
        hi = {c: i for i, c in enumerate(hdr)}
        pcol, poscol = hi.get('Player', 1), hi.get('Pos', 3)
        tkeys = (tags or {}).get(team, {})
        claimed = {}
        for r in rows[1:]:
            full = str(r[pcol])
            parts = full.split()
            if len(parts) < 2:
                continue
            pos = str(r[poscol]).split('-')[0].strip()   # 'PG-SG' -> 'PG'
            if pos not in POSITIONS:
                continue

            key = None
            if tkeys:
                for cand in tag_key_forms(full):
                    if cand in tkeys:
                        key = cand
                        break
            if key is None:
                key = norm(parts[0][0] + '. ' + ' '.join(parts[1:]))

            if key in claimed and claimed[key] != full:
                collisions.append((season, team, key, claimed[key], full))
            claimed[key] = full
            posmap[team][key] = pos
            posmap[team].setdefault(strip_suffix(key), pos)   # alias, never clobber
    wb.close()
    load_positions.collisions = collisions
    return dict(posmap)


def teams_in(season):
    wb = openpyxl.load_workbook(tags_path(season), read_only=True, data_only=True)
    t = sorted({s[:-8] for s in wb.sheetnames
                if s.endswith(' Summary') and not s.startswith('League')})
    wb.close()
    return t


# ---------------------------------------------------------------- scenario data
def extract_scenarios(season):
    """Minute-weighted net rating per scenario, per team, straight from the lineup logs.

    -> {team: {'base_min','base_net','st_short_min',...,'E','B','I','F','L','X','Y'}}
       E = preserve net, B = preserve minutes
       I = pilfer net,   F = pilfer minutes
       L = NRG = E - I           <-- the headline metric
       X = st_deep_net - st_short_net   (4S+1 level gap)
       Y = prsv_net - pilf_net          (bench level gap)
    """
    wb = openpyxl.load_workbook(tags_path(season), data_only=True)
    out = {}
    for team in teams_in(season):
        if team not in wb.sheetnames:
            continue
        agg = defaultdict(lambda: [0.0, 0.0])   # key -> [minutes, minutes*net]
        for r in wb[team].iter_rows(min_row=2, values_only=True):
            if len(r) < 28:
                continue
            key = SCENARIO_KEY.get(r[27])
            mn, net = r[2], r[5]
            if key and isinstance(mn, (int, float)) and isinstance(net, (int, float)):
                agg[key][0] += mn
                agg[key][1] += mn * net

        rec = {}
        for key in SCENARIO_KEY.values():
            m = agg[key][0]
            rec[key + '_min'] = m
            rec[key + '_net'] = (agg[key][1] / m) if m else None

        b = (rec['st_deep_min'] or 0) + (rec['prsv_min'] or 0)
        f = (rec['st_short_min'] or 0) + (rec['pilf_min'] or 0)
        e = ((agg['st_deep'][1] + agg['prsv'][1]) / b) if b else None
        i = ((agg['st_short'][1] + agg['pilf'][1]) / f) if f else None
        rec.update(E=e, B=b, I=i, F=f,
                   L=(e - i) if (e is not None and i is not None) else None)
        rec['X'] = ((rec['st_deep_net'] or 0) - (rec['st_short_net'] or 0)) \
            if rec['st_deep_net'] is not None and rec['st_short_net'] is not None else None
        rec['Y'] = ((rec['prsv_net'] or 0) - (rec['pilf_net'] or 0)) \
            if rec['prsv_net'] is not None and rec['pilf_net'] is not None else None
        rec['tot_min'] = sum((rec[k + '_min'] or 0) for k in SCENARIO_KEY.values())
        out[team] = rec
    wb.close()
    return out


def extract_all_scenarios():
    """-> list of 150 records, each with 's' (season) and 't' (team) plus extract_scenarios keys."""
    data = []
    for s in SEASONS:
        for team, rec in extract_scenarios(s).items():
            row = {'s': s, 't': team}
            row.update(rec)
            data.append(row)
    return data


# ---------------------------------------------------------------- grid / starters
def games_with_non_starters(season, team, tags_for_team, min_non_starters=1):
    """Count games whose actual starting five could NOT be filled by distinct tagged Starters.

    BUG #2: a naive name lookup UNDERCOUNTS teams with same-initial/surname teammates
    (e.g. OKC 2022-23 rostered Jalen AND Jaylin Williams; the Grid shows both as
    'J. Williams'). We therefore run a bipartite max-matching of the five starters
    against the DISTINCT tagged-Starter pool: any starter that cannot be matched to
    its own distinct Starter is a non-Starter.

    Returns (games_played, games_with_>=min_non_starters).
    """
    wb = openpyxl.load_workbook(tags_path(season), read_only=True, data_only=True)
    grid = list(wb[team + ' Grid'].iter_rows(values_only=True))
    wb.close()
    srows = [r for r in grid if r[0] and str(r[0]).startswith('Starter ')][:5]

    starter_keys = [initial_key(nm) for nm, tg in tags_for_team.items() if tg == 'ST']

    def max_match(five_keys):
        adj = [[j for j, pk in enumerate(starter_keys) if sk == pk] for sk in five_keys]
        matched_to = [-1] * len(starter_keys)      # starter slot -> which of the five

        def try_assign(i, seen):
            # NOTE: `seen` must be shared across the whole augmenting-path search for
            # this i. Allocating a fresh `seen` inside the recursion causes infinite
            # mutual recursion when two players share a key (e.g. two 'J. Williams').
            for j in adj[i]:
                if not seen[j]:
                    seen[j] = True
                    if matched_to[j] == -1 or try_assign(matched_to[j], seen):
                        matched_to[j] = i
                        return True
            return False

        return sum(try_assign(i, [False] * len(starter_keys)) for i in range(len(five_keys)))

    played = flagged = 0
    for g in range(1, 83):
        five = [srows[k][g] for k in range(len(srows)) if g < len(srows[k]) and srows[k][g]]
        if not five:
            continue
        played += 1
        n_non = len(five) - max_match([initial_key(x) for x in five])
        if n_non >= min_non_starters:
            flagged += 1
    return played, flagged


def position_starters(season, tags, games, posmap):
    """The Starter with the most starts at each position. -> {team: {POS: normname}}"""
    best = defaultdict(lambda: defaultdict(lambda: (-1, None)))
    for team, players in tags.items():
        for nm, tag in players.items():
            if tag != 'ST':
                continue
            pos = posmap.get(team, {}).get(nm) or posmap.get(team, {}).get(strip_suffix(nm))
            if not pos:
                continue
            gs = games[team][nm][1]
            if gs > best[team][pos][0]:
                best[team][pos] = (gs, nm)
    return {t: {p: best[t][p][1] for p in best[t]} for t in best}


# ---------------------------------------------------------------- Finding 3.9
def by_position(season, apply_rule=True, min_deep_minutes=0):
    """By-position NRG, attributing each preserve/pilfer lineup to every position
    whose starter is OUT of that lineup.

    apply_rule (Keegan's Finding 3.9 methodology rule):
        If a position's starter has NO 'Starters + Deep Bench Player' (4S+1DB) minutes,
        he was never actually replaced by a Deep Bench player. His Short-Bench-Preserved
        minutes are then EXCLUDED and the position is NOT CALCULABLE for that team --
        the team drops out of that position's league weighted average entirely.

    min_deep_minutes: optional floor on 4S+1DB minutes (0 = any amount unlocks the
        position, which is the rule as literally specified; a floor of e.g. 10 prevents
        a single garbage minute from admitting a full ~700-minute preserve bundle).

    -> (league {POS: gap|None}, per_team {team: {POS: gap|None}}, kept {POS: n})
    """
    tags, games = load_tags(season)
    posmap = load_positions(season, tags)
    pstar = position_starters(season, tags, games, posmap)

    wb = openpyxl.load_workbook(tags_path(season), data_only=True)
    acc = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0.0, 0.0])))
    for team in teams_in(season):
        if team not in wb.sheetnames:
            continue
        starters = pstar.get(team, {})
        for r in wb[team].iter_rows(min_row=2, values_only=True):
            if len(r) < 28:
                continue
            key = SCENARIO_KEY.get(r[27])
            if key not in PRESERVE_SCENARIOS | PILFER_SCENARIOS:
                continue
            mn, net = r[2], r[5]
            if not (isinstance(mn, (int, float)) and isinstance(net, (int, float))):
                continue
            on_court = {strip_team(r[i], team) for i in range(13, 18)}
            for pos, starter in starters.items():
                if starter and starter not in on_court:      # this position's starter is OUT
                    cell = acc[team][pos][key]
                    cell[0] += mn
                    cell[1] += mn * net
    wb.close()

    def gap_of(scen):
        deep = scen['st_deep'][0]
        if apply_rule and deep <= min_deep_minutes:
            return None                                   # NOT CALCULABLE
        pre_m = deep + scen['prsv'][0]
        pil_m = scen['st_short'][0] + scen['pilf'][0]
        if not pre_m or not pil_m:
            return None
        pre_n = scen['st_deep'][1] + scen['prsv'][1]
        pil_n = scen['st_short'][1] + scen['pilf'][1]
        return pre_n / pre_m - pil_n / pil_m

    per_team = {t: {p: gap_of(acc[t][p]) for p in POSITIONS} for t in acc}

    pool = defaultdict(lambda: {'pre': [0.0, 0.0], 'pil': [0.0, 0.0]})
    kept = defaultdict(int)
    for team in acc:
        for pos in POSITIONS:
            if per_team[team].get(pos) is None:
                continue
            kept[pos] += 1
            scen = acc[team][pos]
            pool[pos]['pre'][0] += scen['st_deep'][0] + scen['prsv'][0]
            pool[pos]['pre'][1] += scen['st_deep'][1] + scen['prsv'][1]
            pool[pos]['pil'][0] += scen['st_short'][0] + scen['pilf'][0]
            pool[pos]['pil'][1] += scen['st_short'][1] + scen['pilf'][1]

    league = {}
    for pos in POSITIONS:
        pre, pil = pool[pos]['pre'], pool[pos]['pil']
        league[pos] = (pre[1] / pre[0] - pil[1] / pil[0]) if (pre[0] and pil[0]) else None
    return league, per_team, dict(kept)


def by_position_pooled(apply_rule=True, min_deep_minutes=0):
    """Five-season pooled by-position NRG (minute-weighted across all seasons)."""
    pool = defaultdict(lambda: {'pre': [0.0, 0.0], 'pil': [0.0, 0.0]})
    kept = defaultdict(int)
    for season in SEASONS:
        tags, games = load_tags(season)
        posmap = load_positions(season, tags)
        pstar = position_starters(season, tags, games, posmap)
        wb = openpyxl.load_workbook(tags_path(season), data_only=True)
        acc = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0.0, 0.0])))
        for team in teams_in(season):
            if team not in wb.sheetnames:
                continue
            starters = pstar.get(team, {})
            for r in wb[team].iter_rows(min_row=2, values_only=True):
                if len(r) < 28:
                    continue
                key = SCENARIO_KEY.get(r[27])
                if key not in PRESERVE_SCENARIOS | PILFER_SCENARIOS:
                    continue
                mn, net = r[2], r[5]
                if not (isinstance(mn, (int, float)) and isinstance(net, (int, float))):
                    continue
                on_court = {strip_team(r[i], team) for i in range(13, 18)}
                for pos, starter in starters.items():
                    if starter and starter not in on_court:
                        cell = acc[team][pos][key]
                        cell[0] += mn
                        cell[1] += mn * net
        wb.close()
        for team in acc:
            for pos in POSITIONS:
                scen = acc[team][pos]
                deep = scen['st_deep'][0]
                if apply_rule and deep <= min_deep_minutes:
                    continue
                pre_m = deep + scen['prsv'][0]
                pil_m = scen['st_short'][0] + scen['pilf'][0]
                if not pre_m or not pil_m:
                    continue
                kept[pos] += 1
                pool[pos]['pre'][0] += pre_m
                pool[pos]['pre'][1] += scen['st_deep'][1] + scen['prsv'][1]
                pool[pos]['pil'][0] += pil_m
                pool[pos]['pil'][1] += scen['st_short'][1] + scen['pilf'][1]

    league = {}
    for pos in POSITIONS:
        pre, pil = pool[pos]['pre'], pool[pos]['pil']
        league[pos] = (pre[1] / pre[0] - pil[1] / pil[0]) if (pre[0] and pil[0]) else None
    return league, dict(kept)
